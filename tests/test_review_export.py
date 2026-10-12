"""旧审核导出必须支持中文文件名并保持用户隔离。"""
import io
import unittest
from unittest.mock import AsyncMock, patch
from urllib.parse import unquote

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import hardware_auth_middleware
from app.routers.review import router, review_tasks


class ReviewExportTests(unittest.TestCase):
    def test_chinese_filename_exports_valid_workbook(self):
        app = FastAPI()
        app.middleware("http")(hardware_auth_middleware)
        app.include_router(router)
        task = {"status": "completed", "owner_key": "owner", "filename": "合成审核报告.xlsx",
                "result": {"summary": {}, "details": {}, "project_info": {}}}
        with patch.object(review_tasks, "get_task", return_value=task), patch("app.auth._verify", new=AsyncMock(return_value=({"id": "owner"}, True))):
            with TestClient(app) as client:
                response = client.get("/api/v1/review/export/synthetic")
        self.assertEqual(response.status_code, 200)
        header = response.headers["content-disposition"]
        header.encode("ascii")
        self.assertIn("审核报告_合成审核报告.xlsx", unquote(header))
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames[0], "审核概览")
        self.assertEqual(len(workbook.sheetnames), 7)
        with patch.object(review_tasks, "get_task", return_value=task), patch("app.auth._verify", new=AsyncMock(return_value=({"id": "other"}, True))):
            with TestClient(app) as client:
                self.assertEqual(client.get("/api/v1/review/export/synthetic").status_code, 404)
