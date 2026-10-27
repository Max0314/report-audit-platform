import asyncio
import io
import json
import unittest
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.auth import hardware_auth_middleware
from app.routers import cqc
from app.routers.cqc_library import router
from app.services import cqc_library as library
from app.services.cqc_pcb import parse_database, check_pcb_components
from app.services.cqc_report import extract_safety_components, extract_certificate, extract_fields, component_kind


HEADERS = ['名称Name', '型号Type', '规格/材料Specification/Material',
           '商标/制造商Trademark /Manufacture', '生产厂Factory', '认证标准', '认证号']
ROWS = [
    ['印制线路板', 'TEST-01、TEST-02', 'V-0', '甲电子有限公司', '甲电子有限公司', 'GB 4943.1-2022', 'CQC18001185587'],
    ['印制线路板', 'OTHER-01', 'HB', '乙电子有限公司', '乙电子有限公司', 'GB 4943.1-2022', 'CQC21001318856'],
]


def workbook_bytes(rows=None, headers=None, sheet='CQC认证信息'):
    w = Workbook(); s = w.active; s.title = sheet
    s.append(HEADERS if headers is None else headers)
    for row in ROWS if rows is None else rows: s.append(row)
    w.create_sheet('UL（CE）认证信息').append(['ignored'])
    w.create_sheet('BIS认证信息').append(['ignored'])
    output = io.BytesIO(); w.save(output); w.close()
    return output.getvalue()


class MemoryStore:
    def __init__(self): self.objects = {}; self.fail_key = None
    def get_bytes(self, key): return self.objects.get(key)
    def put_bytes(self, key, value, *args):
        if key == self.fail_key: raise OSError('storage failure')
        self.objects[key] = value
    def delete(self, key): self.objects.pop(key, None)


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.store = MemoryStore()
        self.patch = patch('app.services.cqc_library.get_store', return_value=self.store)
        self.patch.start(); self.addCleanup(self.patch.stop)

    def test_publish_deduplicate_update_pin_and_rollback(self):
        content = workbook_bytes()
        first = library.publish(content, 'first.xlsx', 'admin', 0)
        pinned = library.current_snapshot()
        self.assertEqual(len(pinned['records']), 2)
        duplicate = library.publish(content, 'renamed.xlsx', 'admin', 1)
        self.assertTrue(duplicate['unchanged'])
        rows = deepcopy(ROWS); rows[0][1] = 'TEST-03'; rows.pop()
        second = library.publish(workbook_bytes(rows), 'second.xlsx', 'admin', 1)
        self.assertEqual(second['changes'], {'added': 0, 'modified': 1, 'removed': 1, 'unchanged': 0})
        self.assertEqual(pinned['records'][0]['model'], 'TEST-01、TEST-02')
        self.assertEqual(library.current_snapshot()['records'][0]['model'], 'TEST-03')
        library.activate(first['version'], 'admin', 2)
        self.assertEqual(library.current_snapshot(), pinned)
        self.assertEqual(len(library.catalog()['events']), 3)
        self.assertEqual(library.catalog()['revision'], 3)

    def test_invalid_workbook_keeps_current_version(self):
        library.publish(workbook_bytes(), 'valid.xlsx', 'admin', 0)
        original = self.store.objects[library.CATALOG]
        bad = deepcopy(ROWS); bad[1][-1] = bad[0][-1]
        result = library.publish(workbook_bytes(bad), 'bad.xlsx', 'admin', 1)
        self.assertFalse(result['published']); self.assertTrue(result['errors'])
        self.assertEqual(self.store.objects[library.CATALOG], original)

    def test_storage_failure_does_not_change_active_catalog(self):
        first = library.publish(workbook_bytes(), 'valid.xlsx', 'admin', 0)
        self.store.fail_key = library.CATALOG
        with self.assertRaises(OSError):
            library.publish(workbook_bytes(ROWS[:1]), 'next.xlsx', 'admin', 1)
        self.assertEqual(library.catalog()['active'], first['version'])
        self.assertEqual(len(library.catalog()['versions']), 1)

    def test_concurrent_updates_one_wins_and_stale_client_is_rejected(self):
        a, b = workbook_bytes(), workbook_bytes(ROWS[:1])
        def submit(content):
            try: return library.publish(content, 'data.xlsx', 'admin', 0)['published']
            except library.LibraryConflict: return False
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, [a, b]))
        self.assertEqual(sorted(results), [False, True])
        self.assertEqual(library.catalog()['revision'], 1)

    def test_missing_database_or_unavailable_store_never_passes(self):
        self.assertIn('message', library.current_snapshot())
        with patch.object(self.store, 'get_bytes', side_effect=OSError('private detail')):
            snapshot = library.current_snapshot()
        self.assertNotIn('private detail', str(snapshot))
        self.assertEqual(check_pcb_components([], snapshot)['status'], '未完成')


class DatabaseValidationTests(unittest.TestCase):
    def test_exact_template_and_only_cqc_sheet(self):
        records, errors = parse_database(workbook_bytes())
        self.assertFalse(errors); self.assertEqual(len(records), 2)
        self.assertEqual([r['row'] for r in records], [2, 3])

    def test_formula_numeric_certificate_missing_field_and_duplicate_rejected(self):
        for index, value in [(1, '=1+1'), (6, 18001185587), (4, None), (6, ROWS[1][6])]:
            with self.subTest(index=index, value=value):
                rows = deepcopy(ROWS); rows[0][index] = value
                self.assertTrue(parse_database(workbook_bytes(rows))[1])

    def test_missing_sheet_header_empty_and_broken_file(self):
        for content in [b'broken', workbook_bytes(sheet='UL'), workbook_bytes(headers=HEADERS[:-1])]:
            with self.assertRaises(ValueError): parse_database(content)
        self.assertTrue(parse_database(workbook_bytes([]))[1])

    def test_missing_library_records_are_advisory_only(self):
        records, _ = parse_database(workbook_bytes())
        snapshot = {'records': records, 'metadata': {'version': 'v1'}}
        report = {**records[0], 'part_number': '1-1', 'page': 2}
        result = check_pcb_components([report], snapshot)
        self.assertEqual(result['status'], '通过')
        self.assertEqual(result['summary']['missing_records'], 1)
        self.assertTrue(result['warnings'])
        for field in ('model', 'certificate', 'manufacturer', 'factory', 'standard', 'specification'):
            with self.subTest(field=field):
                bad = {**report, field: 'WRONG'}
                self.assertEqual(check_pcb_components([bad], snapshot)['status'], '不通过')
        self.assertEqual(check_pcb_components([], snapshot)['status'], '未完成')

    def test_missing_and_extra_models_and_duplicate_report(self):
        records, _ = parse_database(workbook_bytes())
        snapshot = {'records': records, 'metadata': {}}
        r = {**records[0], 'model': 'TEST-01、UNKNOWN'}
        result = check_pcb_components([r, r], snapshot)
        self.assertEqual(result['checks'][0]['missing_models'], ['TEST-02'])
        self.assertEqual(result['checks'][0]['extra_models'], ['UNKNOWN'])
        self.assertTrue(any(i['kind'] == '重复' for i in result['checks'][1]['issues']))


class ParsingFixTests(unittest.TestCase):
    def test_component_type_is_not_inferred_from_model_substring(self):
        self.assertEqual(component_kind('电源适配器 PCB-TEST 输入：100V'), 'adapter')
        self.assertIsNone(component_kind('PCB-TEST V-0'))

    def test_report_metadata_ends_sample_fields(self):
        self.assertEqual(extract_fields(['完成日期：2026年10月02日\n报告编号：TEST-P-S\n安全关键件清单'])['完成日期']['value'], '2026年10月02日')

    def test_grouped_cross_page_mixed_numbering_and_inherited_name(self):
        pages = ['安全关键件清单 A-P-S\n2. / / 印制板基材（PCB） X V-0\nGB 4943.1-2022 CQC18001185587\n'
                 '/ / Y V-1\nGB 4943.1-2022 CQC21001318856',
                 'A-P-S\n/ / Z HB\nGB 4943.1-2022 CQC24001452089\n'
                 '2-4 / / 印制板基材（PCB） W V-2\nGB 4943.1-2022 CQC10001054669']
        parts = extract_safety_components(pages)
        self.assertEqual([p['part_number'] for p in parts], ['2-1', '2-2', '2-3', '2-4'])
        self.assertEqual([p['model'] for p in parts], ['X', 'Y', 'Z', 'W'])
        self.assertEqual([p['page'] for p in parts], [1, 1, 2, 2])
        self.assertTrue(all(p['certificate'] for p in parts))

    def test_actual_name_overrides_group_number(self):
        parts = extract_safety_components(['安全关键件清单 A-P-S\n1. / / 印制板基材（PCB） X V-0\n'
                                          '甲电子有限公司 甲电子有限公司 GB 4943.1-2022 CQC18001185587'])
        self.assertEqual(parts[0]['name'], '印制板基材（PCB）')
        self.assertEqual(parts[0]['model'], 'X')
        self.assertEqual(parts[0]['manufacturer'], '甲电子有限公司')

    def test_certificate_wrapping_unlabeled_and_ambiguous(self):
        for text in ['CQC18001185587', 'CQC1800118558\n7', '证书号：CQC18001185587']:
            self.assertEqual(extract_certificate('GB 4943.1-2022 ' + text), 'CQC18001185587')
        self.assertIsNone(extract_certificate('GB 4943.1-2022 5000 2026'))
        self.assertIsNone(extract_certificate('CQC18001185587 / CQC21001318856'))
        self.assertEqual(extract_certificate('GB 4943.1-2022\n2023010907572\n089'), '2023010907572089')


class LibraryApiTests(unittest.TestCase):
    def setUp(self):
        self.user = {'id': 'admin', 'capabilities': {'canManageModelConfig': True}}
        self.verify = patch('app.auth._verify', new=AsyncMock(side_effect=lambda *a, **kw: (self.user, True)))
        self.verify.start(); self.addCleanup(self.verify.stop)
        self.store = MemoryStore()
        p = patch('app.services.cqc_library.get_store', return_value=self.store)
        p.start(); self.addCleanup(p.stop)
        app = FastAPI(); app.middleware('http')(hardware_auth_middleware); app.include_router(router)
        self.client = TestClient(app).__enter__(); self.addCleanup(self.client.__exit__, None, None, None)
        self.client.cookies.set('csrf_token', 'test')
        self.headers = {'X-CSRF-Token': 'test'}

    def upload(self, content=None, revision=0, headers=None):
        return self.client.post('/api/v1/cqc-library/upload',
            files={'file': ('data.xlsx', workbook_bytes() if content is None else content)},
            data={'expected_revision': str(revision)}, headers=self.headers if headers is None else headers)

    def test_upload_history_activate_conflict_and_invalid(self):
        first = self.upload().json()
        second = self.upload(workbook_bytes(ROWS[:1]), 1).json()
        self.assertEqual(self.client.get('/api/v1/cqc-library').json()['active'], second['version'])
        self.assertEqual(self.upload(revision=0).status_code, 409)
        reply = self.client.post('/api/v1/cqc-library/activate', headers=self.headers,
                                json={'version': first['version'], 'expected_revision': 2})
        self.assertEqual(reply.status_code, 200)
        self.assertEqual(self.upload(b'broken', 3).status_code, 422)

    def test_anonymous_member_csrf_and_fresh_permission_enforced(self):
        self.assertEqual(self.upload(headers={}).status_code, 403)
        self.user = {'id': 'member', 'capabilities': {}}
        self.assertEqual(self.upload().status_code, 403)
        self.assertEqual(self.client.get('/api/v1/cqc-library').status_code, 403)
        self.user = None
        self.assertEqual(self.upload().status_code, 401)
        self.user = {'id': 'admin', 'capabilities': {'canManageModelConfig': True}}
        with patch('app.auth._verify', new=AsyncMock(side_effect=[(self.user, True), ({'id': 'admin'}, True)])):
            self.assertEqual(self.upload().status_code, 403)

    def test_job_uses_snapshot_captured_before_processing(self):
        first = library.publish(workbook_bytes(), 'first.xlsx', 'admin', 0)
        snapshot = library.current_snapshot()
        library.publish(workbook_bytes(ROWS[:1]), 'next.xlsx', 'admin', 1)
        job = {'task_id': 'a' * 32, 'mode': 'review', 'status': 'processing',
               'owner_key': 'admin', 'source_keys': [], 'pcb_snapshot': snapshot}
        result = {'safety_components': [{**snapshot['records'][0], 'page': 1}], 'warnings': []}
        with patch('app.routers.cqc.persist_upload', new=AsyncMock()), patch('app.routers.cqc.analyze_cqc_pdf', return_value=result), patch('app.routers.cqc.get_store', return_value=self.store):
            asyncio.run(cqc.run_job(job, [('test.pdf', b'%PDF-test')]))
        self.assertEqual(job['status'], 'completed')
        self.assertEqual(job['result']['pcb_check']['database']['version'], first['version'])
        archived = json.loads(self.store.objects[cqc.artifact_key('admin', 'a' * 32)])
        self.assertNotIn('pcb_snapshot', archived)
        self.assertNotIn('pcb_snapshot', job)


if __name__ == '__main__': unittest.main()
