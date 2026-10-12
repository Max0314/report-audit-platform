"""隔离执行的 CQC 作业；临时状态在内存，原文件和结果文件写通 OSS。"""
import asyncio
import hashlib
import io
import json
import time
import uuid
from concurrent.futures import ThreadPoolExecutor

from fastapi import APIRouter, HTTPException, Request, UploadFile, File
from fastapi.responses import Response
from openpyxl import Workbook

from app.auth import current_user_key, require_task_owner
from app.storage import get_store, persist_upload, rollback_objects, safe_filename
from app.services.cqc_report import analyze_cqc_pdf
from app.services.cqc_compare import compare_cqc_pdfs

router = APIRouter(prefix="/api/v1", tags=["CQC"])
MAX_FILE_BYTES = 50 * 1024 * 1024
_executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="cqc")
_jobs: dict[str, dict] = {}
_running: set[asyncio.Task] = set()
_active_owners: set[str] = set()


def artifact_key(owner: str, task_id: str) -> str:
    return f"results/cqc/{hashlib.sha256(owner.encode()).hexdigest()}/{task_id}.json"


def public_job(job: dict) -> dict:
    return {k: v for k, v in job.items() if k not in ("owner_key", "created", "source_keys", "ai")}


async def save_result(job: dict) -> None:
    store = get_store()
    if store:
        await asyncio.to_thread(store.put_bytes, artifact_key(job["owner_key"], job["task_id"]),
                                json.dumps(public_job(job), ensure_ascii=False).encode(), "application/json")


async def run_job(job: dict, files: list[tuple[str, bytes]]) -> None:
    try:
        for index, (name, contents) in enumerate(files):
            job["source_keys"].append(await persist_upload("cqc", job["task_id"], name, contents, index))
        loop = asyncio.get_running_loop()
        if job["mode"] == "review":
            result = await loop.run_in_executor(_executor, analyze_cqc_pdf, files[0][1])
            result["filename"] = files[0][0]
        else:
            result = await loop.run_in_executor(_executor, compare_cqc_pdfs,
                                                files[0][1], files[1][1], files[0][0], files[1][0])
        files.clear()
        job["result"] = result
        job["status"] = "completed"
        await save_result(job)
    except ValueError as exc:
        job.update(status="failed", message=str(exc)[:300])
        await rollback_objects(job["source_keys"])
    except Exception:
        job.update(status="failed", message="处理或 OSS 归档失败，请稍后重试。")
        if "result" not in job:
            await rollback_objects(job["source_keys"])
    finally:
        _active_owners.discard(job["owner_key"])


async def submit(request: Request, mode: str, uploads: list[UploadFile]) -> dict:
    owner = current_user_key(request)
    # 在读入文件之前预留配额，最多两项 CQC 作业，不挤占旧功能的工作线程池。
    if owner in _active_owners or len(_active_owners) >= 2:
        raise HTTPException(429, "CQC 正在处理任务，请稍后重试（每人一项，全局两项）。")
    _active_owners.add(owner)
    try:
        files = []
        for upload in uploads:
            name = safe_filename(upload.filename or "")
            if not name.lower().endswith(".pdf"):
                raise HTTPException(400, "仅支持 PDF 报告")
            contents = await upload.read(MAX_FILE_BYTES + 1)
            if len(contents) > MAX_FILE_BYTES:
                raise HTTPException(413, "单份报告不能超过 50 MB")
            if not contents.startswith(b"%PDF-"):
                raise HTTPException(400, "文件不是有效 PDF")
            files.append((name, contents))
        now = time.time()
        for key in list(_jobs):
            if _jobs[key]["status"] in ("completed", "failed") and (
                    now - _jobs[key]["created"] > 86400 or len(_jobs) >= 40):
                _jobs.pop(key)
        task_id = uuid.uuid4().hex
        job = {"task_id": task_id, "mode": mode, "status": "processing", "owner_key": owner,
               "created": now, "source_keys": []}
        _jobs[task_id] = job
        task = asyncio.create_task(run_job(job, files))
        _running.add(task)
        task.add_done_callback(_running.discard)
        return public_job(job)
    except BaseException:
        _active_owners.discard(owner)
        raise
    finally:
        for upload in uploads:
            await upload.close()


@router.post("/cqc-review/analyze", status_code=202)
async def review(request: Request, file: UploadFile = File(...)):
    return await submit(request, "review", [file])


@router.post("/cqc-compare/analyze", status_code=202)
async def compare(request: Request, old_file: UploadFile = File(...), new_file: UploadFile = File(...)):
    return await submit(request, "compare", [old_file, new_file])


async def load_job(request: Request, task_id: str) -> dict:
    if len(task_id) != 32 or any(c not in "0123456789abcdef" for c in task_id):
        raise HTTPException(404, "任务不存在")
    owner = current_user_key(request)
    if task_id in _jobs:
        require_task_owner(request, _jobs[task_id])
        return public_job(_jobs[task_id])
    store = get_store()
    if store:
        try:
            contents = await asyncio.to_thread(store.get_bytes, artifact_key(owner, task_id))
            if contents:
                job = json.loads(contents)
                if job["status"] != "completed":
                    job["status"] = "completed"
                return public_job(job)
        except Exception:
            raise HTTPException(503, "结果归档暂时不可用") from None
    raise HTTPException(404, "任务不存在或已过期")


@router.get("/cqc/tasks/{task_id}")
async def task_result(request: Request, task_id: str):
    return await load_job(request, task_id)


def export_workbook(job: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CQC 分析结果"
    sheet.append(["项目", "结果"])
    # 保留完整结构和来源，展开为易读路径；去掉图片 base64。
    def append(path, value):
        if isinstance(value, dict):
            for key, child in value.items():
                if not key.endswith("preview"):
                    append(f"{path} / {key}", child)
        elif isinstance(value, list):
            for index, child in enumerate(value, 1):
                append(f"{path} / {index}", child)
        else:
            text = str(value) if value is not None else "未提取到"
            # XLSX 单元格文本上限：拆行，且禁用用户输入公式。
            for start in range(0, max(1, len(text)), 30000):
                sheet.append([path, text[start:start + 30000]])
                sheet.cell(sheet.max_row, 1).data_type = "s"
                sheet.cell(sheet.max_row, 2).data_type = "s"
    append("分析", job)
    sheet.column_dimensions["A"].width = 60
    sheet.column_dimensions["B"].width = 110
    sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get("/cqc/tasks/{task_id}/export")
async def export(request: Request, task_id: str):
    job = await load_job(request, task_id)
    if job["status"] != "completed":
        raise HTTPException(409, "任务尚未完成，暂不能导出")
    content = await asyncio.get_running_loop().run_in_executor(_executor, export_workbook, job)
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="cqc-{task_id}.xlsx"',
                             "Cache-Control": "private, no-store"})
