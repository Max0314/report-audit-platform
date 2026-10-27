"""上传模板校验及 CQC PCB 逐字段核对，不将库内未匹配项判为失败。"""

import re
import unicodedata

from openpyxl import load_workbook


FIELD_LABELS = {
    "name": "关键件名称",
    "model": "型号",
    "specification": "规格/材料",
    "manufacturer": "制造商",
    "factory": "生产厂",
    "standard": "认证标准",
    "certificate": "认证号",
}
HEADER_NAMES = {
    "name": "名称",
    "model": "型号",
    "specification": "规格/材料",
    "manufacturer": "商标/制造商",
    "factory": "生产厂",
    "standard": "认证标准",
    "certificate": "认证号",
}


def normalize_value(value) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"[‐‑‒–—−]", "-", text)
    return re.sub(r"\s+", "", text).upper()


def missing_value(value) -> bool:
    return normalize_value(value) in {"", "/", "-", "未识别", "未提取到"}


def split_models(value) -> list[str]:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"[‐‑‒–—−]", "-", text)
    text = re.sub(r"-\s+(?=[A-Za-z0-9])", "-", text)
    models = {}
    for model in re.split(r"[,、;；\r\n]+", text):
        model = re.sub(r"\s+", "", model)
        if not missing_value(model):
            models.setdefault(normalize_value(model), model)
    return list(models.values())


def model_keys(value) -> set[str]:
    return {normalize_value(model) for model in split_models(value)}


def comparable_value(field: str, value):
    if field == "model":
        return model_keys(value)
    if field == "name" and re.search(r"PCB|印制板|印制线路板|印刷线路板", normalize_value(value)):
        return "PCB"
    if field == "specification":
        return normalize_value(value).removesuffix("级")
    if field == "standard":
        return {item for item in re.split(r"[;,、]+|(?=GB(?:/T)?\d)", normalize_value(value)) if item}
    return normalize_value(value)


def parse_database(contents: bytes) -> tuple[list[dict], list[str]]:
    """只读取上传的 CQC 工作表；校验通过才允许发布整个版本。"""
    import io
    import zipfile
    if len(contents) > 5 * 1024 * 1024:
        raise ValueError("Excel 不能超过 5 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Excel 解压大小或内部文件数超过限制")
            if any(i.filename.lower().endswith("vbaproject.bin") for i in archive.infolist()):
                raise ValueError("不支持包含宏的工作簿")
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError("无法读取 Excel，请上传有效的 .xlsx 文件") from None
    try:
        sheet_name = "CQC认证信息"
        if sheet_name not in workbook.sheetnames:
            raise ValueError("缺少 CQC认证信息 工作表")
        sheet = workbook[sheet_name]
        if sheet.max_row > 1001 or sheet.max_column > 30:
            raise ValueError("CQC 工作表最多 1000 条记录、30 列")
        rows = sheet.iter_rows()
        header = [cell.value for cell in next(rows, ())]
        columns = {}
        for field, label in HEADER_NAMES.items():
            matches = [index for index, value in enumerate(header)
                       if normalize_value(value).startswith(normalize_value(label))]
            if len(matches) != 1:
                raise ValueError(f"表头缺少或重复列：{label}")
            columns[field] = matches[0]
        records, errors, certificates = [], [], set()
        total_chars = 0
        for row_number, cells in enumerate(rows, 2):
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            record = {}
            for field, index in columns.items():
                cell = cells[index]
                value = cell.value
                if cell.data_type in ("f", "e"):
                    errors.append(f"第 {row_number} 行 {FIELD_LABELS[field]} 不允许公式或错误值")
                if value is not None and not isinstance(value, str):
                    errors.append(f"第 {row_number} 行 {FIELD_LABELS[field]} 请使用文本格式，避免编号精度丢失")
                record[field] = str(value).strip() if value is not None else None
                total_chars += len(record[field] or "")
                if total_chars > 1000000:
                    raise ValueError("CQC 认证库文本总量超过限制，请精简清单")
                if len(record[field] or "") > 10000:
                    errors.append(f"第 {row_number} 行 {FIELD_LABELS[field]} 内容过长")
            record.update(row=row_number, sheet=sheet_name)
            empty = [label for field, label in FIELD_LABELS.items() if missing_value(record[field])]
            if empty:
                errors.append(f"第 {row_number} 行缺少：{'、'.join(empty)}")
            if comparable_value("name", record["name"]) != "PCB":
                errors.append(f"第 {row_number} 行不是 PCB 条目")
            certificate = normalize_value(record["certificate"])
            if not re.fullmatch(r"CQC\d{8,16}", certificate):
                errors.append(f"第 {row_number} 行认证号格式错误，应为 CQC 加数字")
            if certificate in certificates:
                errors.append(f"第 {row_number} 行认证号重复，不能唯一匹配")
            certificates.add(certificate)
            records.append(record)
        if not records:
            errors.append("CQC 工作表没有有效记录")
        return records, errors
    finally:
        workbook.close()


def find_candidates(component: dict, records: list[dict]) -> list[dict]:
    certificate = normalize_value(component.get("certificate"))
    if not missing_value(certificate):
        matches = [record for record in records if normalize_value(record["certificate"]) == certificate]
        if matches:
            return matches
    companies = [field for field in ("manufacturer", "factory") if not missing_value(component.get(field))]
    if not companies:
        return []
    matches = [record for record in records
               if all(normalize_value(component[field]) == normalize_value(record[field]) for field in companies)]
    if len(matches) > 1:
        related = [record for record in matches if model_keys(component.get("model")) & model_keys(record["model"])]
        return related or matches
    return matches


def check_pcb_components(components: list[dict], snapshot: dict) -> dict:
    pcb_components = [component for component in components
                      if comparable_value("name", component.get("name")) == "PCB"]
    result = {
        "status": "未完成",
        "mode": "报告条目核对，库内未匹配项仅提示",
        "database": None,
        "checks": [],
        "missing_records": [],
        "database_errors": [],
        "warnings": [],
        "summary": {"total": len(pcb_components), "passed": 0, "failed": 0, "missing_records": 0},
    }
    if not snapshot.get("records"):
        result["warnings"].append(snapshot.get("message", "尚未配置认证库，PCB 核对未完成"))
        return result
    records = snapshot["records"]
    result["database"] = snapshot["metadata"]
    used_certificates = set()
    for component in pcb_components:
        check = {
            "report": component,
            "database_record": None,
            "candidate_rows": [],
            "candidate_sources": [],
            "status": "通过",
            "issues": [],
            "missing_models": [],
            "extra_models": [],
        }
        candidates = find_candidates(component, records)
        check["candidate_rows"] = [record["row"] for record in candidates]
        check["candidate_sources"] = [record.get("source") or {
            "filename": snapshot["metadata"].get("filename"), "sheet": record["sheet"], "row": record["row"],
        } for record in candidates]
        record = candidates[0] if len(candidates) == 1 else None
        check["database_record"] = record
        for field, label in FIELD_LABELS.items():
            report_value = component.get(field)
            if missing_value(report_value):
                check["issues"].append({
                    "field": field, "kind": "缺失", "message": f"报告缺少{label}",
                })
            elif record and comparable_value(field, report_value) != comparable_value(field, record[field]):
                check["issues"].append({
                    "field": field, "kind": "不一致", "message": f"{label}与数据库不一致",
                })
        if record:
            certificate_key = comparable_value("certificate", record["certificate"])
            if certificate_key in used_certificates:
                check["issues"].append({
                    "field": "record", "kind": "重复", "message": "报告重复列出同一数据库 PCB 条目",
                })
            used_certificates.add(certificate_key)
            report_models = model_keys(component.get("model"))
            database_models = model_keys(record["model"])
            check["missing_models"] = [model for model in split_models(record["model"])
                                       if normalize_value(model) not in report_models]
            check["extra_models"] = [model for model in split_models(component.get("model"))
                                     if normalize_value(model) not in database_models]
        else:
            message = "数据库中找不到对应 PCB" if not candidates else "对应多个数据库条目，无法唯一匹配"
            check["issues"].append({"field": "record", "kind": "未匹配", "message": message})
        if check["issues"]:
            check["status"] = (
                "未匹配" if not candidates else
                "匹配不唯一" if record is None else
                "信息缺失" if any(issue["kind"] == "缺失" for issue in check["issues"]) else "不一致"
            )
        result["checks"].append(check)
    result["missing_records"] = [record for record in records
                                 if comparable_value("certificate", record["certificate"]) not in used_certificates]
    result["summary"].update({
        "passed": sum(check["status"] == "通过" for check in result["checks"]),
        "failed": sum(check["status"] != "通过" for check in result["checks"]),
        "missing_records": len(result["missing_records"]),
    })
    if not pcb_components:
        result["warnings"].append("报告未提取到 PCB 条目，无法完成 PCB 核对")
    else:
        result["status"] = "不通过" if result["summary"]["failed"] else "通过"
    if result["missing_records"]:
        result["warnings"].append("数据库中有未在本报告匹配的 PCB，仅作提示，不据此判定不通过")
    if any(check["extra_models"] for check in result["checks"]):
        result["warnings"].append("报告包含数据库未登记型号；请核查数据库版本和证书原件，不能据此认定认证无效")
    return result
