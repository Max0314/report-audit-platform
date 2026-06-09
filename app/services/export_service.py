# -*- coding: utf-8 -*-
"""
硬件研发部测试报告审核平台 - 导出服务
生成审核报告（Excel格式）
"""

import io
from datetime import datetime
from typing import Any, Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core import get_logger

logger = get_logger("export_service")


class ExportService:
    """导出服务"""

    def __init__(self):
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True, size=11)
        self.issue_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
        self.normal_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def generate_review_report(
        self,
        result_data: Dict[str, Any],
        filename: str
    ) -> bytes:
        """生成审核报告（Excel格式）

        Args:
            result_data: 审核结果数据
            filename: 原始文件名

        Returns:
            Excel文件字节流
        """
        wb = Workbook()

        # 1. 概览Sheet
        self._create_summary_sheet(wb, result_data, filename)

        # 2. 测试值逻辑问题Sheet
        self._create_logic_issues_sheet(wb, result_data)

        # 3. 日期问题Sheet
        self._create_date_issues_sheet(wb, result_data)

        # 4. 一致性问题Sheet
        self._create_consistency_issues_sheet(wb, result_data)

        # 5. 完整性问题Sheet
        self._create_completeness_issues_sheet(wb, result_data)

        # 6. AI问题Sheet
        self._create_ai_issues_sheet(wb, result_data)

        # 7. 项目信息Sheet
        self._create_project_info_sheet(wb, result_data)

        # 保存到字节流
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def _create_summary_sheet(self, wb: Workbook, result_data: Dict[str, Any], filename: str):
        """创建概览Sheet"""
        ws = wb.active
        ws.title = "审核概览"

        # 标题
        ws['A1'] = "硬件研发部测试报告审核报告"
        ws['A1'].font = Font(size=16, bold=True, color="4472C4")
        ws.merge_cells('A1:D1')

        # 基本信息
        ws['A3'] = "文件名"
        ws['B3'] = filename
        ws['A4'] = "审核时间"
        ws['B4'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # 统计汇总
        summary = result_data.get("summary", {})
        ws['A6'] = "问题统计"
        ws['A6'].font = Font(bold=True, size=12)

        row = 7
        stats = [
            ("测试值逻辑问题", summary.get("logic_issues", 0)),
            ("日期问题", summary.get("date_issues", 0)),
            ("一致性问题", summary.get("consistency_issues", 0)),
            ("完整性问题", summary.get("completeness_issues", 0)),
            ("AI检查问题", summary.get("ai_issues", 0)),
            ("总问题数", summary.get("total_issues", 0)),
        ]

        for item, count in stats:
            ws[f'A{row}'] = item
            ws[f'B{row}'] = count
            if item == "总问题数":
                ws[f'B{row}'].font = Font(bold=True, color="FF0000")
            row += 1

        # 调整列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15

    def _create_logic_issues_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建测试值逻辑问题Sheet"""
        ws = wb.create_sheet("测试值逻辑问题")

        # 表头
        headers = ["Sheet名称", "行号", "测试项", "实测值", "最小标准", "最大标准", "判定结果", "问题类型", "问题描述"]
        self._write_headers(ws, headers)

        # 数据
        details = result_data.get("details", {})
        value_logic = details.get("value_logic_validation", {})
        issues = value_logic.get("issues", [])

        row = 2
        for issue in issues:
            ws[f'A{row}'] = issue.get("sheet", "")
            ws[f'B{row}'] = issue.get("row", "")
            ws[f'C{row}'] = issue.get("item", "")
            ws[f'D{row}'] = str(issue.get("test_value", ""))
            ws[f'E{row}'] = str(issue.get("min_value", ""))
            ws[f'F{row}'] = str(issue.get("max_value", ""))
            ws[f'G{row}'] = issue.get("result", "")
            ws[f'H{row}'] = issue.get("issue_type", "")
            ws[f'I{row}'] = issue.get("issue", "")

            # 高亮问题行
            for col in range(1, 10):
                ws.cell(row=row, column=col).fill = self.issue_fill

            row += 1

        if not issues:
            ws['A2'] = "未发现测试值逻辑问题"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _create_date_issues_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建日期问题Sheet"""
        ws = wb.create_sheet("日期问题")

        headers = ["Sheet名称", "行号", "列号", "日期值", "错误原因"]
        self._write_headers(ws, headers)

        details = result_data.get("details", {})
        date_validation = details.get("date_validation", {})
        invalid_dates = date_validation.get("invalid", [])

        row = 2
        for item in invalid_dates:
            ws[f'A{row}'] = item.get("sheet", "")
            ws[f'B{row}'] = item.get("row", "")
            ws[f'C{row}'] = item.get("col", "")
            ws[f'D{row}'] = item.get("value", "")
            ws[f'E{row}'] = item.get("error", "")

            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = self.issue_fill

            row += 1

        if not invalid_dates:
            ws['A2'] = "未发现日期问题"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _create_consistency_issues_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建一致性问题Sheet"""
        ws = wb.create_sheet("一致性问题")

        headers = ["字段名称", "存在值", "出现位置", "问题描述"]
        self._write_headers(ws, headers)

        details = result_data.get("details", {})
        consistency = details.get("consistency_validation", {})
        inconsistencies = consistency.get("inconsistencies", [])

        row = 2
        for item in inconsistencies:
            ws[f'A{row}'] = item.get("field", "")
            ws[f'B{row}'] = ", ".join(item.get("values", []))
            ws[f'C{row}'] = ", ".join([str(loc) for loc in item.get("locations", [])])
            ws[f'D{row}'] = item.get("issue", "")

            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = self.issue_fill

            row += 1

        if not inconsistencies:
            ws['A2'] = "未发现一致性问题"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _create_completeness_issues_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建完整性问题Sheet"""
        ws = wb.create_sheet("完整性问题")

        headers = ["类别", "缺失测试项", "状态"]
        self._write_headers(ws, headers)

        details = result_data.get("details", {})
        completeness = details.get("completeness_validation", {})
        missing_items = completeness.get("missing_items", [])

        row = 2
        for item in missing_items:
            ws[f'A{row}'] = item.get("category", "")
            ws[f'B{row}'] = item.get("item", "")
            ws[f'C{row}'] = item.get("status", "")

            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = self.issue_fill

            row += 1

        if not missing_items:
            ws['A2'] = "未发现完整性问题"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _create_ai_issues_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建AI问题Sheet"""
        ws = wb.create_sheet("AI检查问题")

        headers = ["Sheet名称", "位置", "原文", "修正建议", "问题原因"]
        self._write_headers(ws, headers)

        details = result_data.get("details", {})
        ai_result = details.get("ai_result", {})
        issues = ai_result.get("issues", [])

        row = 2
        for item in issues:
            ws[f'A{row}'] = item.get("sheet", "")
            ws[f'B{row}'] = item.get("location", "")
            ws[f'C{row}'] = item.get("original", "")
            ws[f'D{row}'] = item.get("suggestion", "")
            ws[f'E{row}'] = item.get("reason", "")

            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = self.issue_fill

            row += 1

        if not issues:
            ws['A2'] = "未发现AI检查问题"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _create_project_info_sheet(self, wb: Workbook, result_data: Dict[str, Any]):
        """创建项目信息Sheet"""
        ws = wb.create_sheet("项目信息")

        headers = ["字段名称", "提取值"]
        self._write_headers(ws, headers)

        project_info = result_data.get("project_info", {})

        row = 2
        for field, value in project_info.items():
            ws[f'A{row}'] = field
            ws[f'B{row}'] = value or "(未提取到)"
            row += 1

        if not project_info:
            ws['A2'] = "未提取到项目信息"
            ws['A2'].fill = self.normal_fill

        self._adjust_column_width(ws, headers)

    def _write_headers(self, ws, headers: List[str]):
        """写入表头"""
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.border

    def _adjust_column_width(self, ws, headers: List[str]):
        """调整列宽"""
        for col, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = max(15, len(header) + 2)


# 创建导出服务实例
export_service = ExportService()
